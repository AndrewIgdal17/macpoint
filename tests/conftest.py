"""Shared fixtures for MacPoint tests."""

import shutil
import tempfile
import zipfile
from pathlib import Path

import pytest
from pptx import Presentation


@pytest.fixture
def tmp_pptx(tmp_path):
    """Create a minimal .pptx with 3 slides for testing."""
    prs = Presentation()
    layout = prs.slide_layouts[0]
    for i in range(3):
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = f"Slide {i + 1}"
    path = tmp_path / "test.pptx"
    prs.save(str(path))
    return path


@pytest.fixture
def weg_pptx(tmp_path):
    """Create a .pptx from the WEG template for testing."""
    src = Path(__file__).parent.parent / "tests" / "fixtures" / "WEGTemplate.potx"
    if not src.exists():
        # Fall back to vault path for local development
        src = Path(__file__).parent.parent.parent.parent / "Projects/MacPoint/templates/WEGTemplate.potx"
    if not src.exists():
        pytest.skip("WEGTemplate.potx not found")
    dest = tmp_path / "weg_test.pptx"

    _TEMPLATE_CT = b"application/vnd.openxmlformats-officedocument.presentationml.template.main+xml"
    _PRESENTATION_CT = b"application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"

    with zipfile.ZipFile(src, "r") as zin, zipfile.ZipFile(dest, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(_TEMPLATE_CT, _PRESENTATION_CT)
            zout.writestr(item, data)
    return dest
