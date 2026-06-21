"""Tests for add_speaker_notes, slide_snapshot, and manage_slide."""

from pptx import Presentation

from macpoint.backends.pptx_backend import add_speaker_notes, manage_slide, slide_snapshot


def test_add_speaker_notes(tmp_pptx):
    """Should set speaker notes on the specified slide."""
    result = add_speaker_notes(tmp_pptx, 1, "These are my notes")
    assert "slide 1" in result
    assert "18 chars" in result

    prs = Presentation(str(tmp_pptx))
    notes = prs.slides[0].notes_slide.notes_text_frame.text
    assert notes == "These are my notes"


def test_add_speaker_notes_overwrites(tmp_pptx):
    """Calling twice should overwrite, not append."""
    add_speaker_notes(tmp_pptx, 1, "First notes")
    add_speaker_notes(tmp_pptx, 1, "Second notes")

    prs = Presentation(str(tmp_pptx))
    notes = prs.slides[0].notes_slide.notes_text_frame.text
    assert notes == "Second notes"


def test_slide_snapshot_basic(tmp_pptx):
    """Should return shape/placeholder info for a slide."""
    result = slide_snapshot(tmp_pptx, 1)
    assert "Slide 1/3" in result
    assert "Slide 1" in result  # title text
    assert "Shapes" in result
    assert "Placeholders" in result


def test_slide_snapshot_default_slide(tmp_pptx):
    """With no slide_number, should default to slide 1."""
    result = slide_snapshot(tmp_pptx, None)
    assert "Slide 1/3" in result


def test_manage_slide_delete(tmp_pptx):
    """Should delete a slide and reduce count."""
    result = manage_slide(tmp_pptx, "delete", 2)
    assert "Deleted slide 2" in result
    assert "2 slides" in result

    prs = Presentation(str(tmp_pptx))
    assert len(prs.slides) == 2


def test_manage_slide_move(tmp_pptx):
    """Should move a slide to a new position."""
    result = manage_slide(tmp_pptx, "move", 1, target_position=3)
    assert "Moved slide 1 to position 3" in result

    prs = Presentation(str(tmp_pptx))
    assert prs.slides[2].shapes.title.text == "Slide 1"
    assert prs.slides[0].shapes.title.text == "Slide 2"


def test_manage_slide_duplicate(tmp_pptx):
    """Should duplicate a slide (appended to end)."""
    result = manage_slide(tmp_pptx, "duplicate", 1)
    assert "Duplicated slide 1" in result
    assert "4 slides" in result

    prs = Presentation(str(tmp_pptx))
    assert len(prs.slides) == 4
    # The duplicate should have the same title text
    assert prs.slides[3].shapes.title.text == "Slide 1"


def test_manage_slide_delete_last_fails(tmp_pptx):
    """Cannot delete the only remaining slide."""
    manage_slide(tmp_pptx, "delete", 1)
    manage_slide(tmp_pptx, "delete", 1)
    # Now only 1 slide remains
    try:
        manage_slide(tmp_pptx, "delete", 1)
        assert False, "Should have raised"
    except ValueError as e:
        assert "last remaining" in str(e)
